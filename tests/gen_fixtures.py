"""Generate test fixture documents (Chinese + English) under tests/fixtures/."""
import os
import csv
from pathlib import Path

OUT = Path(__file__).parent / "fixtures"
OUT.mkdir(parents=True, exist_ok=True)

# ---------- TXT ----------
(OUT / "sample.txt").write_text(
    "这是一个中文 TXT 测试文档。\nHello, world!\n第二行：apple 苹果 / banana 香蕉。\n",
    encoding="utf-8",
)

# ---------- Markdown ----------
(OUT / "sample.md").write_text(
    "# 中文标题\n\n这是 **Markdown** 测试。\n\n- 项目一\n- 项目二\n",
    encoding="utf-8",
)

# ---------- CSV ----------
with open(OUT / "sample.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.writer(f)
    w.writerow(["姓名", "年龄", "城市"])
    w.writerow(["张三", "28", "北京"])
    w.writerow(["李四", "31", "上海"])
    w.writerow(["Alice", "25", "Seattle"])

# ---------- HTML ----------
(OUT / "sample.html").write_text(
    """<!doctype html><html><head><title>测试</title>
    <style>body{color:red}</style></head>
    <body><h1>中文 HTML 标题</h1>
    <p>这是一段含 <b>加粗</b> 与 <i>斜体</i> 的文字。</p>
    <script>console.log('skip me')</script>
    <ul><li>苹果</li><li>香蕉</li></ul>
    </body></html>""",
    encoding="utf-8",
)

# ---------- JSON ----------
(OUT / "sample.json").write_text(
    '{"name": "张三", "age": 28, "tags": ["开发者", "Rust"]}\n',
    encoding="utf-8",
)

# ---------- DOCX ----------
from docx import Document
from docx.shared import Inches
import io, struct, zlib

def _tiny_png(color=(0x33, 0x66, 0x99)):
    """Build a 16×16 solid-color PNG entirely in-process (no external deps)."""
    w = h = 16
    raw = b"".join(b"\x00" + bytes(color) * w for _ in range(h))
    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(
            ">I", zlib.crc32(tag + data) & 0xFFFFFFFF
        )
    sig = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0)
    idat = zlib.compress(raw, 9)
    return sig + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")

png_bytes = _tiny_png()
(OUT / "_tiny.png").write_bytes(png_bytes)

doc = Document()
doc.add_heading("中文 Word 测试文档", level=1)
doc.add_paragraph("这是第一段。包含中文与 English mixed content。")
doc.add_picture(str(OUT / "_tiny.png"), width=Inches(1))
doc.add_heading("二级标题", level=2)
doc.add_paragraph("项目符号列表：")
doc.add_paragraph("苹果", style="List Bullet")
doc.add_paragraph("香蕉", style="List Bullet")
table = doc.add_table(rows=2, cols=2)
table.rows[0].cells[0].text = "姓名"
table.rows[0].cells[1].text = "城市"
table.rows[1].cells[0].text = "张三"
table.rows[1].cells[1].text = "北京"
doc.save(OUT / "sample.docx")

# ---------- DOCX (multi-page with explicit page breaks) ----------
big = Document()
big.add_heading("中文 Word 大文档", level=1)
big.add_paragraph("封面：这是一份多页测试文档。包含 5 页正文，每页含中英文混排。")
big.add_page_break()
for page in range(1, 6):
    big.add_heading(f"第 {page} 页 / Page {page}", level=2)
    for _ in range(3):
        big.add_paragraph(
            f"段落示例：Rust 是一门系统编程语言。Page {page} content. "
            f"中文标点：你好，世界！This is a longer paragraph with "
            f"mixed 中英文 to test extraction quality. " * 2
        )
    if page != 5:
        big.add_page_break()
big.save(OUT / "big.docx")

# ---------- XLSX ----------
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "员工"
ws.append(["姓名", "年龄", "城市"])
ws.append(["张三", 28, "北京"])
ws.append(["李四", 31, "上海"])
ws2 = wb.create_sheet("Numbers")
ws2.append(["A", "B", "C"])
ws2.append([1, 2, 3])
wb.save(OUT / "sample.xlsx")

# ---------- PPTX ----------
from pptx import Presentation
from pptx.util import Inches
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "中文 PPT 测试"
slide.placeholders[1].text = "副标题：第一张幻灯片"
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "议程"
tf = slide2.placeholders[1].text_frame
tf.text = "项目一：背景"
tf.add_paragraph().text = "项目二：方案"
tf.add_paragraph().text = "项目三：总结"
prs.save(OUT / "sample.pptx")

# ---------- PDF ----------
# WeasyPrint produces a properly-structured PDF with a complete ToUnicode CMap,
# which is what real-world PDFs (Word/browser/Acrobat) look like and what
# pure-Rust `pdf-extract` can decode well.
from weasyprint import HTML
HTML(string="""<html><head><meta charset='utf-8'>
<style>
  body { font-family: 'WenQuanYi Zen Hei','Droid Sans Fallback',sans-serif; font-size: 13pt; }
  h1 { font-size: 20pt; }
  .page-break { page-break-before: always; }
</style></head>
<body>
  <h1>中文 PDF 测试文档</h1>
  <p>第一行：你好，世界。Hello, world!</p>
  <p>第二行：苹果 / 香蕉 / 橙子。</p>
  <div class='page-break'></div>
  <h1>第二页 - Page 2</h1>
  <p>Rust 是一门系统编程语言。</p>
  <p>It mixes 中英文 in the same sentence.</p>
</body></html>""").write_pdf(str(OUT / "sample.pdf"))

print("Fixtures written to:", OUT)
for p in sorted(OUT.iterdir()):
    print(" -", p.name, p.stat().st_size, "bytes")
